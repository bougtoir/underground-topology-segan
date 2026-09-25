"""Insert figures inline into manuscript_draft.docx -> manuscript_inline.docx."""
import os, copy
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
doc = Document(f"{ROOT}/manuscript/manuscript_draft.docx")

insertions = [  # (anchor_text, [fig files])
    ("Mean TLD values", ["fig1_study_zones.png", "fig2_topology_example.png"]),
    ("negative 5th percentile", ["fig3_ltp.png", "fig4_tornado.png"]),
    ("(Figure 5)", ["fig5_dg.png"]),
]
caps = {
    "fig1_study_zones.png": "Figure 1. Study zones: road graph (grey) and snapped building demand nodes (blue).",
    "fig2_topology_example.png": "Figure 2. Example feeder cluster: S1 like-for-like replicate vs S2 optimized layout.",
    "fig3_ltp.png": "Figure 3. Life-cycle topology profit by zone (base case vs Monte-Carlo mean).",
    "fig4_tornado.png": "Figure 4. One-at-a-time sensitivity (tornado) of LTP, Sano zone.",
    "fig5_dg.png": "Figure 5. NPC of S2 vs S3 (rooftop-PV extension) by zone.",
}

def insert_after(par, figs):
    anchor = par._p
    for f in figs:
        pic_par = doc.add_paragraph()
        pic_par.add_run().add_picture(f"{ROOT}/figures/{f}", width=Inches(5.8))
        cap_par = doc.add_paragraph(caps[f])
        cap_par.runs[0].font.size = None
        anchor.addnext(cap_par._p)
        anchor.addnext(pic_par._p)
        anchor = cap_par._p

done = set()
for par in doc.paragraphs:
    for anchor_text, figs in insertions:
        if anchor_text in par.text and anchor_text not in done:
            insert_after(par, figs)
            done.add(anchor_text)
print("inserted:", done)
doc.save(f"{ROOT}/manuscript/manuscript_inline.docx")
print("saved")
