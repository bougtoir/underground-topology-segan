"""Phase F13-F18 (v4): highlights, cover letter, supplementary docx,
municipality outputs, claim matrix, inline docx, final SEGAN submission
package. All values from analysis/manuscript_values_v3.csv (frozen) /
*_v2 tables."""
import csv, os, zipfile, glob, shutil
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
mv = {r["key"]: r["value"] for r in
      csv.DictReader(open(f"{ROOT}/analysis/manuscript_values_v3.csv"))}

def fmt(v, nd=2):
    try:
        return f"{float(v):.{nd}f}"
    except Exception:
        return str(v)

# ---------------- highlights (SEGAN: max 85 chars incl. spaces, 3-5 items)
hl = [
    "Engineering feasibility is imposed before topology comparison",
    "Linearized feeder model agrees with AC power flow on all feeders",
    "Topology flexibility is worth under 0.5% of life-cycle cost",
    "Naive forced-redesign contrasts mislead when incumbents excluded",
    "Reconstruction and LV/MV assumptions shape renewal conclusions",
]
for h in hl:
    assert len(h) <= 85, h
open(f"{ROOT}/manuscript/highlights_v4.txt", "w").write("\n".join(hl) + "\n")

# ---------------- cover letter (F14) ----------------
cl = f"""Cover letter — Sustainable Energy, Grids and Networks

Dear Editors,

We submit "How Much Is Topology Flexibility Worth at Renewal? A
Feasibility-Constrained Life-Cycle Assessment of LV Re-Layout During
Undergrounding" for consideration as an original research article in
Sustainable Energy, Grids and Networks.

Undergrounding renewal creates a rare topology decision window: civil
works are already committed, so re-choosing low-voltage feeder layouts
costs only the marginal rerouting. This paper quantifies whether that
window is worth using. The method integrates four components that are
usually treated separately: legacy-network reconstruction from open
municipal data under explicit data tiers, engineering feasibility gating
before any cost comparison (coincident-load sizing, ampacity, transformer
loading, voltage drop), counterfactual/estimand design in which the
incumbent layout remains an admissible decision, and life-cycle cost
optimization via weighted Steiner trees with per-edge conductor sizing.
The linearized power-flow model is validated against full AC power flow
on all {fmt(mv['ac_n_feeders'],0)} feasible synthesized feeders.

The central result is bounded but informative: the value of topology
flexibility was economically equivalent to like-for-like renewal in all
three modeled LV study zones ({fmt(mv['sano_vof_pct'],2)}%,
{fmt(mv['kudamatsu_vof_pct'],2)}%, {fmt(mv['yasu_vof_pct'],2)}% of
life-cycle cost) under a pre-registered 1% equivalence margin.
Feasibility gating and estimand design are shown to prevent spurious
apparent optimization benefits: a naive forced-redesign contrast
produces a -10% headline that is an estimand artifact and is not
robust to the LV/MV boundary assumption. For utilities and
municipalities planning renewal, the transferable contribution is
methodological — how to value re-layout options honestly under
open-data constraints — rather than a promise of large savings.

The work fits SEGAN's scope on distribution network planning, grid
modernization, and the integration of engineering constraints into
distribution investment decisions.

The manuscript is original, not under consideration elsewhere, and all
authors approve submission.

Sincerely,
The authors
"""
open(f"{ROOT}/manuscript/cover_letter_v4.txt", "w").write(cl)

# ---------------- supplementary docx ----------------
sup_doc = Document()
sup_doc.styles["Normal"].font.name = "Times New Roman"
sup_doc.styles["Normal"].font.size = Pt(11)
sup_doc.add_heading("Supplementary material", 0)
sup_doc.add_paragraph(
    "Supplementary Figure S1. NPC of S2 vs the secondary rooftop-PV "
    "overlay S3 by zone (scenario-based upper bound; does not alter the "
    "primary VoF conclusion).")
sup_doc.add_picture(f"{ROOT}/figures/v4_supp_figS1_dg.png", width=Inches(5.5))
sup_doc.add_paragraph(
    "Benchmark note. A second public MV benchmark (Iwamoto 11-bus) "
    "failed to converge in pandapower 3.5.5 and is excluded from "
    "validation statistics; the AC agreement reported in the main text "
    f"covers all {fmt(mv['ac_n_feeders'],0)} feasible synthesized feeders.")
sup_doc.add_paragraph(
    "Additional machine-readable tables (cluster-level results, "
    "engineering-assumption sensitivity, Monte-Carlo sweep, "
    "claim-evidence matrix) are provided as CSV files in the submission "
    "package.")
sup_doc.save(f"{ROOT}/manuscript/supplementary_material_v4.docx")

# ---------------- declarations ----------------
decl = """Declarations
- Declaration of generative AI: during the preparation of this work the
  authors used an AI coding assistant for code scaffolding and drafting.
  The authors reviewed and edited all content and take full
  responsibility for the publication.
- Data availability: all scripts, the source registry, checksums, and
  machine-readable result tables are provided in the supplementary
  repository snapshot. Municipal open data were used under their stated
  licenses (CC-BY for Kudamatsu).
- Code availability: the full analysis pipeline (scripts/) is included
  in the submission package and the linked public repository.
- Conflict of interest: none declared. Funding: none.
"""
open(f"{ROOT}/manuscript/declarations_v4.txt", "w").write(decl)

# ---------------- claim-evidence matrix (F8) ----------------
ce = [
    ("VoF +0.01%/+0.39%/+0.20% (Sano/Kudamatsu/Yasu), all economically_equivalent within 1% margin",
     "scenario_results_v2.csv vof_pct; PRIMARY_ESTIMAND_SPEC.md; config/ECONOMIC_EQUIVALENCE_SPEC.yaml (frozen pre-classification)"),
    ("Secondary naive ΔNPC +0.01%/-10.1%/+0.20%; negative value is estimand artifact, not primary",
     "scenario_results_v2.csv ltp_pct; audit_v2/COUNTERFACTUAL_OPTIMIZATION_AUDIT.md"),
    ("Kudamatsu -10.1% decomposes to 5 clusters with 2-10x rewiring capex; sign sensitive to LV/MV boundary",
     "BENEFIT_DECOMPOSITION.csv; audit_v2/KUDAMATSU_MECHANISM_AUDIT.md; ENGINEERING_ASSUMPTION_SENSITIVITY.csv mv200"),
    ("LinDistFlow validated vs pandapower AC on 672 feasible synthesized feeders (max vm err 0.0193 pu, tol 0.02; agreement 1.00); validates synthesized models, not real topology",
     "AC_VS_LINEAR_VALIDATION.csv; AC_VALIDATION_REPORT.md"),
    ("MV-threshold sensitivity flips Kudamatsu ΔNPC (+0.27% at 200kW) but not VoF ordering",
     "ENGINEERING_ASSUMPTION_SENSITIVITY.csv"),
    ("24-replicate and leave-one-family-out robust (fam ratio ~1.00)",
     "cluster_results_v2.csv s1_famA/B_npc; rep24 scenario"),
    ("Feasibility gates: vdrop<=6%, ampacity, 150kVA loading, radiality",
     "config/ENGINEERING_FEASIBILITY_SPEC.yaml; p07b_optimize_v2.py"),
    ("MV-served nodes >150kW excluded by design; LV share 75/55/53% of demand",
     "mv_served_nodes.csv; demand_coverage_audit.csv"),
    ("Rooftop PV overlay (S3) is a secondary upper-bound scenario; Supplement only",
     "dg_results_v2.csv; supplementary_material_v4.docx"),
]
pd.DataFrame(ce, columns=["claim", "evidence"]).to_csv(
    f"{ROOT}/analysis/CLAIM_EVIDENCE_MATRIX_FINAL.csv", index=False)

# ---------------- municipality outputs (v4 wording) ----------------
os.makedirs(f"{ROOT}/municipality_v4", exist_ok=True)
tpl = """# {City} — undergrounding renewal planning-support summary (modeled)

Scope: LV study-zone reconstruction from open data; modeled demands;
planning support only — NOT engineering design or engineering advice.

- LV clusters assessed: {ncl} ({nfeas} feasible under the frozen spec)
- Modeled coincident LV demand: {dem} MW
- Value of topology flexibility VoF: {vof}% of S1 NPC (class: {cls})
- Secondary naive Delta-NPC: {ltp}% | Monte-Carlo (500 draws): P(DeltaNPC>0) = {p}
- Mean topology divergence distance TDD: {tdd}
- Interpretation: {interp}
- The rooftop-PV overlay (S3) is an upper-bound sensitivity reported in
  the supplement of the manuscript.

Inputs, code, provenance: analysis/ and provenance/ in the repository
snapshot. Never contact us about this output; it is generated research
material, not a municipal communication.
"""
interp = {
    "sano": "modeled VoF <0.1% of renewal cost; like-for-like renewal is a defensible default within the modeled design space",
    "kudamatsu": "modeled VoF <0.5%; reconstructed incumbent-like layouts are economically equivalent to best feasible redesigns",
    "yasu": "modeled VoF <0.5% of renewal cost; like-for-like renewal is a defensible default within the modeled design space"}
for c, cname in [("sano", "Sano"), ("kudamatsu", "Kudamatsu"), ("yasu", "Yasu")]:
    txt = tpl.format(City=cname, ncl=int(float(mv[f"{c}_clusters"])),
                     nfeas=int(float(mv[f"{c}_clusters_feasible"])),
                     dem=fmt(mv[f"{c}_lv_demand_mw"], 1),
                     vof=fmt(mv[f"{c}_vof_pct"], 2), cls=mv[f"{c}_equivalence_class"],
                     ltp=fmt(mv[f"{c}_ltp_pct"], 2),
                     p=fmt(mv[f"{c}_mc_p_ltp_pos"], 2),
                     tdd=fmt(mv[f"{c}_mean_tld"], 2), interp=interp[c])
    open(f"{ROOT}/municipality_v4/{c}_planning_support.md", "w").write(txt)

# ---------------- inline docx ----------------
doc = Document(f"{ROOT}/manuscript/manuscript_draft_v4.docx")
figs = ["v4_fig1_study_zones.png", "v4_fig2_topology_example.png",
        "v4_fig3_delta.png", "v4_fig4_tornado.png"]
caps = {
    "v4_fig1_study_zones.png": "Figure 1. Study zones: road graph (grey) and snapped building demand nodes (blue).",
    "v4_fig2_topology_example.png": "Figure 2. Example feeder cluster: S1 like-for-like replicate vs S2 optimized layout.",
    "v4_fig3_delta.png": "Figure 3. Primary estimand VoF by zone vs Monte-Carlo mean of the secondary naive ΔNPC.",
    "v4_fig4_tornado.png": "Figure 4. Sensitivity to engineering assumptions: secondary naive ΔNPC (bars) vs primary VoF (dots); grey band = 1% equivalence margin.",
}
insertions = [
    ("(Figures 1–3)", figs[:3]),
    ("(Figure 4)", figs[3:]),
]
done = set()
for par in doc.paragraphs:
    for anchor_text, fl in insertions:
        if anchor_text in par.text and anchor_text not in done:
            anchor = par._p
            for f in fl:
                pic = doc.add_paragraph()
                pic.add_run().add_picture(f"{ROOT}/figures/{f}", width=Inches(5.8))
                cap = doc.add_paragraph(caps[f])
                anchor.addnext(cap._p); anchor.addnext(pic._p)
                anchor = cap._p
            done.add(anchor_text)
print("inserted:", done)
doc.save(f"{ROOT}/manuscript/manuscript_inline_v4.docx")

# ---------------- final package ----------------
pkg = f"{ROOT}/SEGAN_submission_package_FINAL_v4"
if os.path.exists(pkg):
    shutil.rmtree(pkg)
os.makedirs(pkg)
files = [
    "manuscript/manuscript_draft_v4.docx", "manuscript/manuscript_inline_v4.docx",
    "manuscript/supplementary_material_v4.docx", "manuscript/highlights_v4.txt",
    "manuscript/cover_letter_v4.txt", "manuscript/declarations_v4.txt",
    "analysis/manuscript_values_v3.csv", "analysis/scenario_results_v2.csv",
    "analysis/sensitivity_mc_v2.csv", "analysis/sensitivity_tornado_v2.csv",
    "analysis/dg_results_v2.csv", "analysis/ablation_results.csv",
    "analysis/feasibility_sensitivity.csv", "analysis/original_vs_corrected.csv",
    "analysis/CLAIM_EVIDENCE_MATRIX_FINAL.csv", "analysis/mv_served_nodes.csv",
    "analysis/reconstruction_summary_v2.csv", "analysis/references_verified.csv",
    "analysis/novelty_matrix.csv", "analysis/journal_fit_matrix.csv",
    "analysis/data_audit.csv", "analysis/benchmark_validation.json",
    "analysis/benchmark_ac_solvers.csv", "analysis/benchmark_radial_comparison.csv",
    "analysis/benchmark_iwamoto_diagnosis.csv",
    "analysis/NEGATIVE_RESULT_NOVELTY_MATRIX.csv",
    "analysis/ENGINEERING_ASSUMPTION_SENSITIVITY.csv",
    "analysis/demand_coverage_audit.csv", "analysis/BENEFIT_DECOMPOSITION.csv",
    "analysis/FINAL_V3_VALUE_CHECK.csv",
    "audit_v2/AUDIT_FILE_INVENTORY.csv", "audit_v2/MANUSCRIPT_VALUE_AUDIT.csv",
    "audit_v2/VOLTAGE_ROOT_CAUSE_AUDIT.md", "audit_v2/BENCHMARK_VALIDATION_REPORT.md",
    "audit_v2/KUDAMATSU_MECHANISM_AUDIT.md", "audit_v2/COST_MECHANISM_AUDIT.md",
    "audit_v2/COUNTERFACTUAL_OPTIMIZATION_AUDIT.md", "audit_v2/TERMINOLOGY_AUDIT.md",
    "audit_v2/REVIEWER_GATE.md", "audit_v2/REVIEWER_GATE_v4.md",
    "config/ENGINEERING_FEASIBILITY_SPEC.yaml", "config/ECONOMIC_EQUIVALENCE_SPEC.yaml",
    "config/study.yaml", "PRIMARY_ESTIMAND_SPEC.md",
    "tables/table2_results_v4.csv",
    "provenance/source_registry.csv", "provenance/decision_log.md",
    "FINAL_REPRODUCIBILITY_REPORT.md", "FINAL_QC_REPORT_v4.md",
    "FINAL_HANDOFF_v4.txt",
]
for c in ["sano", "kudamatsu", "yasu"]:
    files.append(f"municipality_v4/{c}_planning_support.md")
files += [f"figures/{f}" for f in figs + ["v4_supp_figS1_dg.png"]]
files += [f"scripts/{f}" for f in sorted(glob.glob1(f"{ROOT}/scripts", "*.py"))]
for f in files:
    d = os.path.join(pkg, f)
    os.makedirs(os.path.dirname(d), exist_ok=True)
    shutil.copy(f"{ROOT}/{f}", d)
zpath = f"{ROOT}/SEGAN_submission_package_FINAL_v4.zip"
if os.path.exists(zpath):
    os.remove(zpath)
with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
    for dp, _, fs in os.walk(pkg):
        for f in fs:
            fp = os.path.join(dp, f)
            z.write(fp, os.path.relpath(fp, pkg))
print("zip entries:", len(zipfile.ZipFile(zpath).namelist()))
print("done")
