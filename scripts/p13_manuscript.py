"""Build SEGAN submission documents from the frozen value ledger."""
import csv
import os

from docx import Document
from docx.oxml import parse_xml
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
values = {
    row["key"]: row["value"]
    for row in csv.DictReader(open(
        f"{ROOT}/FINAL_MANUSCRIPT_VALUES.csv", encoding="utf-8"))
}
references = {
    row["doi"]: row
    for row in csv.DictReader(open(
        f"{ROOT}/analysis/references_verified.csv", encoding="utf-8"))
}
CITIES = ["sano", "kudamatsu", "yasu"]
NAMES = {"sano": "Sano", "kudamatsu": "Kudamatsu", "yasu": "Yasu"}


def v(key):
    return float(values[key])


def f(key, digits=2):
    return f"{v(key):.{digits}f}"


def bn(value):
    return f"{float(value) / 1e9:.2f}"


def pct_change(city):
    return f"{v(city + '_peak_loss_change_pct'):+.1f}"


doc = Document()
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(11)


def equation(text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.font.name = "Cambria Math"


doi = {
    "baran": "10.1109/61.25627",
    "distflow": "10.1109/61.19265",
    "pandapower": "10.1109/TPWRS.2018.2829021",
    "steiner1": "10.1002/net.3230220105",
    "steiner2": "10.1007/BF00288961",
    "osm": "10.1109/MPRV.2008.80",
    "conductor": "10.1109/59.43199",
    "underground": "10.1016/j.epsr.2022.108804",
    "linopf": "10.1109/PSCC.2014.7038399",
    "planning": "10.1109/PESMG.2013.6672615",
    "autonomous": "10.1016/j.apenergy.2023.121522",
    "geographic": "10.35833/mpce.2024.000884",
}
order = []


def cite(paragraph, key):
    identifier = doi[key]
    if identifier not in order:
        order.append(identifier)
    run = paragraph.add_run(str(order.index(identifier) + 1))
    run.font.superscript = True


def add_table(title, headers, body):
    caption = doc.add_paragraph(title, style="Caption")
    caption.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row in body:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
    for row in table.rows[:-1]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    for row in table.rows:
        row._tr.get_or_add_trPr().append(parse_xml(
            '<w:cantSplit xmlns:w="http://schemas.openxmlformats.org/'
            'wordprocessingml/2006/main"/>'))
    return table


doc.add_heading(
    "Engineering-feasible screening of underground distribution renewal: "
    "separating routing-policy and conductor-sizing effects", 0)
doc.add_paragraph(
    "Original research article for Sustainable Energy, Grids and Networks")
doc.add_paragraph("[Authors and affiliations to be completed locally]")
doc.add_paragraph("Corresponding author: [to be completed locally]")

doc.add_heading("Abstract", 1)
doc.add_paragraph(
    "Underground distribution renewal can preserve a reconstructed route or "
    "use the construction window to screen alternative routes and equipment. "
    "External analysts often lack utility topology and loading records, making "
    "the defensible scope of such screening uncertain. We develop a reproducible "
    "geospatial experiment with explicit evidence labels, reconstructed radial "
    f"route ensembles, a {f('mv_nominal_kv',1)}-kV/"
    f"{f('lv_nominal_v',0)}–{f('service_nominal_v',0)}-V electrical hierarchy, "
    "catalogued "
    "cables and transformers, iterative active/reactive loss propagation, and "
    "mandatory connectivity, voltage, thermal and power-balance gates. Three "
    "Japanese municipal study zones are treated as synthetic screening cases, "
    "not observed utility networks. A fixed-conductor minimum-construction "
    "policy changes net present cost by "
    f"{f('sano_routing_policy_savings_pct', 1)}%, "
    f"{f('kudamatsu_routing_policy_savings_pct', 1)}% and "
    f"{f('yasu_routing_policy_savings_pct', 1)}%, whereas best-found combined "
    "routing and bounded conductor screening changes it by "
    f"{f('sano_combined_screening_savings_pct', 1)}%, "
    f"{f('kudamatsu_combined_screening_savings_pct', 1)}% and "
    f"{f('yasu_combined_screening_savings_pct', 1)}%. A paired ablation "
    "separates routing-policy and conductor-sizing contributions. Conditional "
    "scenario sampling varies reconstruction, candidate policy and cost/load "
    "assumptions while holding topology and equipment fixed within each draw. "
    "The results show that route reconsideration and conductor sizing must not "
    "be conflated. They support transparent early-stage screening under stated "
    "assumptions, not claims about actual municipal feeders, project savings or "
    "globally optimal designs.")
doc.add_paragraph(
    "Keywords: distribution planning; undergrounding; conductor sizing; "
    "Steiner tree; voltage feasibility; reconstruction uncertainty")

doc.add_heading("1. Introduction", 1)
p = doc.add_paragraph(
    "Renewal of overhead distribution assets creates a choice between "
    "reproducing an incumbent corridor underground and screening alternative "
    "radial routes and equipment. Underground construction can reduce exposure "
    "to some external hazards, but project-specific capital and repair costs "
    "make screening assumptions consequential. ")
cite(p, "underground")
p.add_run(
    " Distribution planning has long represented reconfiguration, expansion "
    "and conductor allocation as coupled decisions. ")
cite(p, "baran")
p.add_run(" ")
cite(p, "planning")
doc.add_paragraph(
    "Recent methods address topology planning with learning-based expansion "
    "optimization and jointly screen geographically constrained feeder routing "
    "and conductor sizing. ")
p = doc.paragraphs[-1]
cite(p, "autonomous")
p.add_run(" ")
cite(p, "geographic")
p.add_run(
    " Those methods motivate the decision structure here, but they assume "
    "decision-ready network models. The present contribution is instead an "
    "evidence-bound workflow for sparse public data; it is not a claim of a "
    "superior optimization algorithm.")

doc.add_paragraph(
    "The difficulty is greater when the analyst has roads and building "
    "footprints but not utility feeder connectivity, conductor records, phase "
    "assignments, switches, protection zones or customer loads. In that setting, "
    "a reconstructed graph is a modeled object. Treating it as an observed "
    "network, or treating a joint routing-and-sizing contrast as a pure topology "
    "effect, would overstate the evidence.")

doc.add_paragraph(
    "This study contributes a forensic and reproducible screening framework "
    "that (i) labels observed, reconstructed, modeled and assumed quantities; "
    "(ii) rejects electrically infeasible candidates; (iii) distinguishes a "
    "fixed-conductor routing-policy contrast from a combined routing-and-sizing "
    "contrast; (iv) compares the radial approximation quantitatively with AC "
    "power flow; and (v) reports uncertainty as conditional scenario sampling "
    "rather than calibrated project probability. The analysis does not claim "
    "global optimization or construction readiness.")

doc.add_heading("2. Methods", 1)
doc.add_heading("2.1 Evidence boundary and modeled demand", 2)
p = doc.add_paragraph(
    "Roads and building footprints were archived from OpenStreetMap. ")
cite(p, "osm")
p.add_run(
    " Buildings define modeled demand geography, not observed customer load, "
    "phase connection or transformer assignment. Municipal security-light "
    "inventories provide location evidence in Sano and Kudamatsu. Sano has "
    f"{int(v('sano_rows')):,} records; {int(v('sano_duplicate_extra_records')):,} "
    "are additional records sharing an identifier, so record coordinates are "
    "not asserted to be exact pole coordinates. Kudamatsu has "
    f"{int(v('kudamatsu_rows')):,} security-light identifiers, which are not "
    "pole identifiers. Yasu contributes no municipal pole or light inventory. "
    "None of the three cases observes feeder connectivity or equipment.")

doc.add_paragraph(
    "Each building's modeled peak demand equals projected footprint area "
    "multiplied by the reported number of floors, or two assumed floors when "
    "that tag is absent, and an assumed class-specific density of 20–60 W/m². "
    "The values are summed after snapping building centroids to road nodes. No "
    "empirical load calibration, customer coincidence factor or temporal "
    "diversity model is applied; the resulting aggregate is a deliberately "
    "conservative coincident screening load, not an observed municipal peak.")

doc.add_paragraph(
    "Positive modeled demand nodes were clustered only within connected road-"
    "graph components, with singleton components retained. The Sano, Kudamatsu "
    f"and Yasu cases contain {int(v('sano_clusters'))}, "
    f"{int(v('kudamatsu_clusters'))} and {int(v('yasu_clusters'))} clusters "
    "and modeled coincident peak demands of "
    f"{v('sano_total_demand_kw')/1000:.2f}, "
    f"{v('kudamatsu_total_demand_kw')/1000:.2f} and "
    f"{v('yasu_total_demand_kw')/1000:.2f} MW, respectively (Figure 1; "
    "Table 1).")
add_table(
    "Table 1. Evidence and modeled-study-zone summary.",
    ["Municipality", "Municipal evidence tier", "Observed records",
     "Modeled clusters", "Modeled demand (MW)"],
    [[
        NAMES[city], values[city + "_evidence_tier"],
        int(v(city + "_rows")), int(v(city + "_clusters")),
        f"{v(city + '_total_demand_kw')/1000:.2f}",
    ] for city in CITIES])

doc.add_heading("2.2 Reconstruction and scenario definitions", 2)
doc.add_paragraph(
    "Each non-singleton cluster uses two reconstruction families: jittered, "
    "anchor-biased Steiner candidates and shuffled incremental cheapest-"
    "attachment candidates. Each family contains "
    f"{int(v('reconstruction_replicates_per_family'))} replicates. Steiner "
    "construction is a graph heuristic, not proof of a globally optimal utility "
    "layout.")
p = doc.add_paragraph("The graph approximations follow established work. ")
cite(p, "steiner1")
p.add_run(" ")
cite(p, "steiner2")

doc.add_paragraph(
    "S0 is a reconstructed legacy route. S1 places reference conductors on "
    "each S0 tree. S1.5 is a fixed-conductor minimum-construction Steiner "
    "policy. S2 is the best-found member of a declared weighted-candidate set "
    "after bounded one-edge-at-a-time economic conductor search. Figure 2 "
    "shows representative modeled routes. Route difference is Jaccard edge-set "
    "divergence:")
equation(
    "D_edge = 1 - |E_S1 intersect E_S2| / |E_S1 union E_S2|")
doc.add_paragraph(
    "The primary fixed-conductor routing-policy estimand is "
    "[NPC(S1)−NPC(S1.5)]/NPC(S1). The combined screening estimand is "
    "[NPC(S1)−NPC(S2)]/NPC(S1). Neither is a likelihood, causal treatment "
    "effect or calibrated project saving.")

doc.add_heading("2.3 Electrical hierarchy and feasibility", 2)
doc.add_paragraph(
    f"The model uses a {f('mv_nominal_kv', 1)}-kV balanced three-phase primary, "
    f"catalogued single-phase {f('mv_nominal_kv',1)}-kV/"
    f"{int(v('lv_nominal_v'))}–"
    f"{int(v('service_nominal_v'))}-V transformers and a single-phase "
    f"{int(v('lv_nominal_v'))}-V outer-conductor secondary equivalent at power "
    f"factor {f('power_factor', 2)}. Sano is evaluated at "
    f"{f('sano_frequency_hz', 0)} Hz and Kudamatsu and Yasu at "
    f"{f('kudamatsu_frequency_hz', 0)} Hz. Each road-snapped demand site has "
    "an explicit transformer. Loads are divided into planning units no larger "
    f"than {f('max_lv_unit_kw', 0)} kW, each with a "
    f"{f('lv_stub_length_m', 0)}-m equivalent LV stub. These are assumptions "
    "required because customer-level secondary geometry is unavailable.")

doc.add_paragraph(
    "Active and reactive customer power plus line losses are propagated "
    "upstream to convergence. For three-phase MV, I=S/(√3V), voltage change is "
    "√3I(R cosφ+X sinφ), and loss is 3I²R. For the single-phase LV equivalent, "
    "I=S/V, voltage change is 2I(R cosφ+X sinφ), and loss is 2I²R. Transformer "
    "no-load and load losses are added upstream. Source input must equal "
    "customer demand plus modeled losses within the frozen tolerance.")
p = doc.add_paragraph(
    "The approximation is related to radial distribution-flow models. ")
cite(p, "distflow")
p.add_run(" ")
cite(p, "linopf")

doc.add_paragraph(
    f"MV, transformer and LV voltage-drop gates are {f('mv_drop_max_pct',1)}%, "
    f"{f('transformer_drop_max_pct',1)}% and {f('lv_drop_max_pct',1)}%. "
    f"Equivalent customer voltage must remain within {f('customer_min_pu',3)}–"
    f"{f('customer_max_pu',3)} pu. Cable and transformer loading cannot exceed "
    f"{f('maximum_loading_pct',0)}%; sizing targets "
    f"{f('design_loading_pct',0)}%. Cable ampacity is multiplied by an assumed "
    f"{f('buried_ampacity_derating',2)} buried-duct factor. Connectivity, "
    "radiality and complete demand service are mandatory. No voltage clipping, "
    "limit relaxation or exclusion of failed clusters is permitted.")

doc.add_paragraph(
    f"Each connected cluster has an independent local "
    f"{f('source_voltage_pu',1)}-pu source. Shared "
    "substations, upstream feeders, switching, protection coordination and "
    "inter-cluster constraints are outside the model boundary. Consequently, "
    "the gates establish internal feasibility of the screening model, not "
    "feasibility of an actual utility system.")

doc.add_heading("2.4 Static-horizon objective and attribution", 2)
doc.add_paragraph(
    "The objective includes trench/duct, conductor and transformer capital "
    "cost; annual O&M; variable line and transformer load loss; and transformer "
    "no-load loss. It discounts a static representative annual loss estimate "
    f"over {int(v('horizon_years'))} years at "
    f"{100*v('discount_rate'):.0f}%. It does not include replacement timing, "
    "failure rates, salvage, reliability, outage damage or construction "
    "staging, and is therefore described as a static-horizon NPC screen rather "
    "than a complete life-cycle model.")
equation(
    "NPC = C_cap + [C_cap f_OM + E_loss p_e] "
    "sum_(t=1)^h (1+r)^(-t)")
doc.add_paragraph(
    "Conductor unit prices scale from the reference area with an assumed "
    f"exponent of {f('conductor_area_cost_exponent',2)}. A bounded local search "
    "accepts feasible one-step conductor upsizes that reduce NPC until no such "
    "move remains. The procedure is best-found and locally screened; it is not "
    "a mixed-integer global-optimality certificate.")

p = doc.add_paragraph(
    "This installed-cost-versus-loss trade-off is consistent with classical "
    "economic conductor selection. ")
cite(p, "conductor")
doc.add_paragraph(
    "A paired 2×2 policy ablation compares a representative reconstructed "
    "legacy route and the minimum-construction route, each with reference and "
    "bounded-search conductors. Shapley averaging attributes the paired "
    "contrast to routing policy and conductor sizing. A second accounting "
    "decomposition reports trench, conductor, transformer and loss components. "
    "These are model-policy attributions, not causal marginal effects.")

doc.add_heading("2.5 Validation and conditional scenario sampling", 2)
p = doc.add_paragraph(
    "The balanced radial calculation was compared with pandapower AC power "
    "flow ")
cite(p, "pandapower")
p.add_run(
    " on three radial feeders extracted from the CIGRE LV benchmark. "
    f"Predeclared tolerances were {f('benchmark_voltage_tolerance_pu',3)} pu "
    "for minimum-voltage error and "
    f"{100*v('benchmark_loss_relative_tolerance'):.0f}% for relative line-loss "
    "error. The comparison validates only the balanced radial-line "
    "approximation. It does not validate the municipal reconstruction, local-"
    f"source boundary, transformer placement, "
    f"{f('service_nominal_v',0)}-V service abstraction, "
    "unbalanced operation, protection or contingencies. The separate "
    "case11_iwamoto network failed to converge under all attempted solvers and "
    "is excluded from validation claims.")

doc.add_paragraph(
    f"Conditional scenario sampling uses {int(v('scenario_sampling_draws')):,} "
    "draws per municipality. Each draw coherently selects one reconstruction "
    "family and replicate across the city and one S2 candidate-policy weight, "
    "then varies construction, conductor and transformer costs; conductor "
    "area-price exponent; electricity value; O&M; discount rate; horizon; "
    "load factor; and demand. Topology and equipment are not reoptimized or "
    "resized within a draw. Reported intervals and positive-draw frequencies "
    "therefore describe the declared frozen-design scenario sampler, not "
    "calibrated probabilities or confidence intervals for real projects.")

doc.add_heading("3. Results", 1)
doc.add_heading("3.1 Feasibility and benchmark agreement", 2)
doc.add_paragraph(
    "All clusters entering S1, S1.5 and S2 summaries passed the frozen model "
    "gates. Across the three cases, maximum modeled MV drop was "
    f"{max(v(c + '_max_mv_vdrop_pct') for c in CITIES):.2f}% and minimum "
    "equivalent customer voltage was "
    f"{min(v(c + '_min_customer_voltage_pu') for c in CITIES):.4f} pu "
    "(Table 2). These replace the quarantined "
    f"{f('quarantined_original_min_voltage_drop_pct',0)}–"
    f"{f('quarantined_original_max_voltage_drop_pct',0)}% primary-drop results, "
    "which combined aggregation and network-structure defects.")
add_table(
    "Table 2. Feasible base-case screening results.",
    ["Municipality", "S1 NPC (bn JPY)", "S1.5 NPC (bn JPY)",
     "S2 NPC (bn JPY)", "Routing-policy savings (%)",
     "Combined savings (%)", "Max MV drop (%)", "Min customer voltage (pu)"],
    [[
        NAMES[city], bn(v(city + "_s1_npc")), bn(v(city + "_s15_npc")),
        bn(v(city + "_s2_npc")),
        f"{v(city + '_routing_policy_savings_pct'):.2f}",
        f"{v(city + '_combined_screening_savings_pct'):.2f}",
        f"{v(city + '_max_mv_vdrop_pct'):.3f}",
        f"{v(city + '_min_customer_voltage_pu'):.4f}",
    ] for city in CITIES])
doc.add_paragraph(
    f"All {int(v('benchmark_feeders_validated'))} benchmark feeders passed. "
    "Maximum minimum-voltage error was "
    f"{f('benchmark_max_voltage_error_pu',6)} pu and maximum relative line-loss "
    f"error was {100*v('benchmark_max_loss_relative_error'):.2f}%.")

doc.add_heading("3.2 Routing-policy and combined screening contrasts", 2)
doc.add_paragraph(
    "The fixed-conductor S1-to-S1.5 routing-policy savings are "
    f"{f('sano_routing_policy_savings_pct',2)}%, "
    f"{f('kudamatsu_routing_policy_savings_pct',2)}% and "
    f"{f('yasu_routing_policy_savings_pct',2)}%. The combined best-found "
    "S1-to-S2 screening savings are "
    f"{f('sano_combined_screening_savings_pct',2)}%, "
    f"{f('kudamatsu_combined_screening_savings_pct',2)}% and "
    f"{f('yasu_combined_screening_savings_pct',2)}%, corresponding to "
    f"{bn(v('sano_delta_npc'))}, {bn(v('kudamatsu_delta_npc'))} and "
    f"{bn(v('yasu_delta_npc'))} billion JPY (Table 2; Figure 3). The large "
    "difference between the two contrasts demonstrates why joint routing and "
    "conductor effects cannot be labeled as topology value.")

doc.add_paragraph(
    "S2 peak loss changes relative to S1 are "
    f"{pct_change('sano')}%, {pct_change('kudamatsu')}% and "
    f"{pct_change('yasu')}%. The objective prices losses but does not minimize "
    "them independently; feasible lower-cost conductors can increase loss. The "
    "repaired results therefore do not support a general claim that route and "
    "equipment screening halves peak loss.")

doc.add_heading("3.3 Paired attribution and scenario sampling", 2)
doc.add_paragraph(
    "In the paired representative-route ablation, routing-policy Shapley "
    "contributions are "
    f"{bn(v('sano_routing_policy_shapley_savings_jpy'))}, "
    f"{bn(v('kudamatsu_routing_policy_shapley_savings_jpy'))} and "
    f"{bn(v('yasu_routing_policy_shapley_savings_jpy'))} billion JPY; "
    "conductor contributions are "
    f"{bn(v('sano_conductor_shapley_savings_jpy'))}, "
    f"{bn(v('kudamatsu_conductor_shapley_savings_jpy'))} and "
    f"{bn(v('yasu_conductor_shapley_savings_jpy'))} billion JPY "
    "(Figure 4; Table 3). The attribution is specific to this paired policy "
    "contrast and candidate ensemble.")
add_table(
    "Table 3. Paired routing-policy/conductor attribution.",
    ["Municipality", "Routing-policy Shapley (bn JPY)",
     "Conductor Shapley (bn JPY)", "Interaction (bn JPY)"],
    [[
        NAMES[city],
        bn(v(city + "_routing_policy_shapley_savings_jpy")),
        bn(v(city + "_conductor_shapley_savings_jpy")),
        bn(v(city + "_routing_conductor_interaction_jpy")),
    ] for city in CITIES])

doc.add_paragraph(
    "Conditional scenario-sampling means are "
    f"{f('sano_sampling_combined_savings_pct_mean',2)}%, "
    f"{f('kudamatsu_sampling_combined_savings_pct_mean',2)}% and "
    f"{f('yasu_sampling_combined_savings_pct_mean',2)}%. The 2.5th–97.5th "
    "percentile ranges are "
    f"[{f('sano_sampling_combined_savings_pct_p2_5',2)}, "
    f"{f('sano_sampling_combined_savings_pct_p97_5',2)}]%, "
    f"[{f('kudamatsu_sampling_combined_savings_pct_p2_5',2)}, "
    f"{f('kudamatsu_sampling_combined_savings_pct_p97_5',2)}]% and "
    f"[{f('yasu_sampling_combined_savings_pct_p2_5',2)}, "
    f"{f('yasu_sampling_combined_savings_pct_p97_5',2)}]% (Table 4). "
    "Topology and equipment remain frozen within every draw; these ranges must "
    "not be interpreted as calibrated project uncertainty.")
add_table(
    "Table 4. Conditional scenario sampling with topology and equipment "
    "frozen within each draw.",
    ["Municipality", "Mean combined savings (%)", "2.5th percentile (%)",
     "97.5th percentile (%)", "Positive-draw frequency"],
    [[
        NAMES[city],
        f"{v(city + '_sampling_combined_savings_pct_mean'):.2f}",
        f"{v(city + '_sampling_combined_savings_pct_p2_5'):.2f}",
        f"{v(city + '_sampling_combined_savings_pct_p97_5'):.2f}",
        f"{v(city + '_sampling_positive_draw_frequency'):.3f}",
    ] for city in CITIES])

doc.add_heading("4. Discussion", 1)
doc.add_paragraph(
    "The principal result is a separation of estimands. Under common reference "
    "conductors, choosing a minimum-construction routing policy produces a small "
    "NPC contrast. The larger combined S1–S2 contrast additionally reflects "
    "bounded conductor sizing. This distinction is scientifically and "
    "practically important: a planner should first test whether a like-for-like "
    "renewal case embeds an uneconomic equipment standard, then evaluate whether "
    "the incremental route effect justifies survey, easement and implementation "
    "costs omitted from the graph screen.")

doc.add_paragraph(
    "The direct within-study baselines also delimit the optimization claim. S1 "
    "tests reconstructed legacy policies; S1.5 tests a minimum-construction "
    "Steiner policy with fixed conductors; the paired ablation tests routing and "
    "conductor decisions in both orders; and S2 searches a finite weighted "
    "candidate set. These comparisons demonstrate improvement over declared "
    "baselines, not superiority over all distribution-planning algorithms.")

doc.add_paragraph(
    "Candidate-route diversity is limited: "
    f"{int(v('candidate_clusters_one_unique_route'))} of "
    f"{int(v('candidate_clusters_total'))} clusters have only one unique route "
    "in the declared weight grid, and "
    f"{int(v('candidate_clusters_selected_weight_zero'))} select the zero-"
    "weight boundary. This is a negative diagnostic, not evidence of a broadly "
    "effective topology optimizer. It confines the route result to the tested "
    "heuristics and makes the fixed-conductor S1.5 comparison the more "
    "transferable routing-policy baseline.")

doc.add_paragraph(
    "The municipal evidence does not support claims about actual feeder routes "
    "or bankable savings. Demand is synthetic, routes are reconstructed, each "
    "cluster has an independent idealized source, transformer sites are modeled, "
    "and secondary circuits are reduced to equivalent stubs. The "
    f"{f('buried_ampacity_derating',2)} ampacity "
    "factor and conductor price curve are assumptions rather than installation-"
    "specific ratings or quotations. Shared-substation capacity, phase imbalance, "
    "switching, reliability, faults, protection, harmonics, transients, "
    "construction access and permitting are omitted.")

doc.add_paragraph(
    "The uncertainty analysis also has a deliberate boundary. It coherently "
    "samples reconstruction and candidate policies and varies a broad parameter "
    "set, but it does not redesign equipment after demand or price changes. "
    "Intervals may therefore understate adaptation available to a planner and "
    "cannot be interpreted as posterior or frequentist uncertainty about a "
    "real project. The within-ensemble prefix-stability diagnostic likewise "
    "does not establish independent-seed algorithmic convergence.")

doc.add_paragraph(
    "A rooftop-PV scenario from the canonical package is excluded from the "
    "headline analysis. Building footprints do not establish roof suitability, "
    "ownership, hourly generation, hosting capacity, curtailment, export value "
    "or reinforcement requirements. A defensible extension requires time-series "
    "load and irradiance and unbalanced hosting-capacity analysis.")

doc.add_paragraph(
    "Future work should use utility topology, equipment and interval-load data; "
    "model shared upstream assets; jointly redesign under each scenario; add "
    "reliability and resilience value; obtain local construction bids; and "
    "compare heuristic candidates with exact or bounded mathematical programs "
    "on tractable instances.")

doc.add_heading("5. Conclusions", 1)
doc.add_paragraph(
    "A repaired, feasibility-gated analysis retains positive combined screening "
    "contrasts in three synthetic municipal cases, but the fixed-conductor "
    "routing-policy effects are much smaller. The evidence supports a "
    "reproducible early-stage method for separating routing and conductor "
    "decisions under explicit assumptions. It does not identify observed "
    "utility topology, certify global optimality, estimate calibrated project "
    "probability or replace detailed engineering.")

doc.add_heading("Declarations", 1)
doc.add_paragraph(
    "Funding: none. Declaration of competing interests: none declared. "
    "Ethics approval: not applicable; the study uses "
    "public infrastructure and geographic data and contains no human participants.")
doc.add_paragraph(
    "Data and code availability: archived raw public-data snapshots, retrieval "
    "metadata, checksums, processing scripts, machine-readable results and "
    "reproduction instructions accompany the submission package.")
doc.add_paragraph(
    "Declaration of generative AI and AI-assisted technologies: during "
    "preparation, the authors used an AI coding assistant for code review, "
    "reproducibility checking and language drafting. The authors reviewed and "
    "edited all outputs and take responsibility for the content.")

doc.add_heading("Figure captions", 1)
captions = [
    "Figure 1. Municipal study zones. Grey lines are archived OpenStreetMap "
    "roads and blue points are road-snapped modeled building-demand nodes. "
    "Neither layer reveals utility connectivity.",
    "Figure 2. Representative reconstructed legacy route and best-found S2 "
    "candidate. Both are modeled screening objects, not observed utility routes.",
    "Figure 3. Base-case combined screening contrast and conditional scenario-"
    "sampling percentiles. Topology and equipment are frozen within each draw.",
    "Figure 4. Paired Shapley attribution to routing policy and conductor "
    "sizing. Values are policy-screening attributions, not causal effects.",
]
for caption in captions:
    doc.add_paragraph(caption)

doc.add_heading("References", 1)
for index, identifier in enumerate(order, 1):
    ref = references[identifier]
    doc.add_paragraph(
        f"{index}. {ref['authors']} {ref['title']} {ref['container']} "
        f"{ref['year']};{ref['volume']}:{ref['pages']}. "
        f"https://doi.org/{identifier}")

os.makedirs(f"{ROOT}/manuscript", exist_ok=True)
manuscript_path = f"{ROOT}/manuscript/manuscript_corrected.docx"
doc.save(manuscript_path)

highlights = [
    "• Separates routing-policy effects from combined routing and sizing.",
    "• Rejects candidates that fail voltage, thermal, radiality or balance gates.",
    "• Validates the balanced radial approximation against three CIGRE feeders.",
    "• Samples reconstruction and policy choices coherently at city level.",
    "• Frames municipal cases as synthetic screens, not observed utility grids.",
]
for line in highlights:
    if len(line.replace("• ", "")) > 85:
        raise RuntimeError(f"highlight exceeds 85 characters: {line}")
with open(f"{ROOT}/manuscript/highlights_corrected.txt", "w", encoding="utf-8") as handle:
    handle.write("\n".join(highlights) + "\n")

supp = Document()
supp.styles["Normal"].font.name = "Times New Roman"
supp.styles["Normal"].font.size = Pt(11)
supp.add_heading("Supplementary methods and reproducibility notes", 0)
supp.add_heading("S1. Evidence classes", 1)
supp.add_paragraph(
    "Observed quantities are direct fields in archived source files. "
    "Reconstructed quantities are graph objects inferred from observed "
    "geography. Modeled quantities are outputs of declared algorithms. Assumed "
    "quantities are engineering or economic inputs without case-specific "
    "measurement. Sensitivity quantities are values sampled over declared ranges.")
supp.add_heading("S2. Engineering gates", 1)
for text in [
        "Connectivity, radiality and complete demand service are mandatory.",
        f"MV, transformer and LV drop limits are {f('mv_drop_max_pct',1)}%, "
        f"{f('transformer_drop_max_pct',1)}% and {f('lv_drop_max_pct',1)}%.",
        f"Customer voltage must remain within {f('customer_min_pu',3)}–"
        f"{f('customer_max_pu',3)} pu.",
        f"Cable and transformer loading cannot exceed "
        f"{f('maximum_loading_pct',0)}%.",
        "Source power must equal customer demand plus modeled losses.",
    ]:
    supp.add_paragraph(text, style="List Bullet")
supp.add_heading("S3. Interpretation limits", 1)
supp.add_paragraph(
    "The cases are synthetic geospatial screens with local idealized sources. "
    "Conditional scenario sampling freezes topology and equipment within each "
    "draw. Prefix stability uses one candidate ensemble and is not independent-"
    "seed algorithmic convergence. The conductor search is bounded and local.")
supp.add_heading("S4. Reproducibility", 1)
supp.add_paragraph(
    "The package contains raw-data checksums and registry, frozen engineering "
    "and study configuration, executable scripts, unit tests, machine-readable "
    "results, the frozen manuscript-value ledger and a clean-room protocol.")
supp.save(f"{ROOT}/manuscript/supplementary_material_corrected.docx")

print("wrote", manuscript_path)
