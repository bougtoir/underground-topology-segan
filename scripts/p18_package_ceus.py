"""CEUS submission package: cover letter, highlights, declarations,
data/code availability, supplement, inline docx, readme, zip.
Frozen values only."""
import csv, os, zipfile, glob, shutil
from docx import Document
from docx.shared import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
mv = {r["key"]: r["value"] for r in
      csv.DictReader(open(f"{ROOT}/analysis/manuscript_values_v3.csv"))}

def fmt(v, nd=2):
    try:
        return f"{float(v):.{nd}f}"
    except Exception:
        return str(v)

# ---------------- highlights (Elsevier: <=85 chars, 3-5) ----------------
hl = [
    "Hidden low-voltage networks reconstructed from open data in three tiers",
    "Feasibility gating precedes any cost comparison of network layouts",
    "Topology flexibility at renewal is worth under 0.5% of life-cycle cost",
    "Residual re-layout value is size-structured, not morphology-driven",
    "Incumbent-admissible estimand prevents inflated redesign savings",
]
for h in hl:
    assert len(h) <= 85, h
open(f"{ROOT}/manuscript/highlights_ceus.txt", "w").write("\n".join(hl) + "\n")

# ---------------- cover letter ----------------
cl = f"""Cover letter — Computers, Environment and Urban Systems

Dear Editor,

I am pleased to submit "The Value of Topology Flexibility at
Infrastructure Renewal: A Feasibility-Gated Geospatial Assessment of
Low-Voltage Re-Layout During Undergrounding" for consideration as a
research article in Computers, Environment and Urban Systems.

Cities renewing aging overhead lines by undergrounding routinely pay
for the corridor but never price the design option it opens: is
re-drawing the hidden local network worth anything once the trench is
committed? This paper turns that question into a computable geospatial
decision problem and answers it end to end. I reconstruct legacy
low-voltage networks from open municipal and OpenStreetMap data in
three observability tiers, impose engineering feasibility gates before
any cost comparison, validate the linearized electrical model against
AC power flow on all {fmt(mv['ac_n_feeders'],0)} feasible synthesized
feeders, and define an incumbent-admissible estimand — the value of
topology flexibility — that keeps the inherited layout inside the
choice set rather than assuming it away.

The manuscript speaks to an existing CEUS conversation: rehabilitation
decision-support for buried utility networks is a journal-native theme
since UtilNets (Hadzilacos et al., 2000), and CEUS has published
network inference from incomplete surface evidence for wastewater
(Chahinian et al., 2019) and BIM/geospatial integration under
incomplete semantics (Gilbert et al., 2021). This paper adds what
those lines still lack: a priced, engineering-feasible decision layer
on reconstructed urban utility topology.

The contribution to CEUS's community is, I believe, threefold. First,
a decision layer on top of the geospatial reconstruction literature:
the reconstructed network is treated not as an answer but as the
choice set over which a renewal option is priced. Second, an honest
bounded result: the option value is under 0.5% of life-cycle cost in
all three study zones, and I show that naive incumbent-excluding
contrasts manufacture double-digit artifacts that disappear under
correct estimand design. Third, a spatial finding reported honestly:
the residual re-layout value concentrates in a small minority of
larger feasible clusters — but every apparent morphology correlate
collapses under cluster-size adjustment, so I report it as an
exploratory screening diagnostic rather than a morphology rule. The method, data provenance, and full reproducibility
package accompany the submission.

The manuscript is original, is not under consideration elsewhere, and
has been approved for submission.

Sincerely,
Tatsuki Onishi
"""
open(f"{ROOT}/manuscript/cover_letter_ceus.txt", "w").write(cl)

# ---------------- declarations ----------------
decl = """Declarations
- Declaration of generative AI: during the preparation of this work the
  author used an AI coding assistant for code scaffolding and drafting.
  The author reviewed and edited all content and takes full
  responsibility for the publication.
- Conflict of interest: none declared. Funding: none.
- Data availability: see data_code_availability_ceus.txt.
"""
open(f"{ROOT}/manuscript/declarations_ceus.txt", "w").write(decl)

dca = """Data and code availability
All analysis scripts, the source registry, checksums, and
machine-readable result tables are provided in the supplementary
repository snapshot at github.com/bougtoir/underground-topology-segan
(reproduce-only scope). Municipal open data were used under their
stated licenses (CC-BY for Kudamatsu); OpenStreetMap data are used
under ODbL. Processed intermediate tables are included in the
submission package; raw municipal extracts are retained locally per
license terms. Figure 2 map layers: geoBoundaries Japan ADM1
(https://www.geoboundaries.org/, CC-BY-4.0) and Natural Earth 10m land
(https://www.naturalearthdata.com/, public domain); plotting script and
raw downloads are included in the repository snapshot.
"""
open(f"{ROOT}/manuscript/data_code_availability_ceus.txt", "w").write(dca)

# ---------------- supplement ----------------
sup_doc = Document()
sup_doc.styles["Normal"].font.name = "Times New Roman"
sup_doc.styles["Normal"].font.size = Pt(11)
sup_doc.add_heading("Supplementary material", 0)
sup_doc.add_paragraph(
    "Supplementary Figure S1. NPC of S2 vs the secondary rooftop-PV "
    "overlay S3 by zone (scenario-based upper bound; does not alter the "
    "primary VoF conclusion).")
sup_doc.add_picture(f"{ROOT}/figures/v4_supp_figS1_dg.png", width=Inches(5.5))
sup_doc.add_paragraph(
    "Benchmark note. A second public MV benchmark (Iwamoto 11-bus) "
    "failed to converge in pandapower 3.5.5 and is excluded from "
    "validation statistics; the AC agreement reported in the main text "
    f"covers all {fmt(mv['ac_n_feeders'],0)} feasible synthesized feeders.")
import csv as _csv
sup_doc.add_paragraph(
    "Supplementary Table S2. Within-zone Moran's I of cluster VoF "
    "(inverse-distance weights on demand-terminal centroids; "
    "200-permutation p-values). Near-zero values indicate negligible "
    "spatial clustering of VoF within zones.")
_sd = list(_csv.DictReader(open(f"{ROOT}/analysis/CEUS_SPATIAL_DIAGNOSTICS.csv")))
_t = sup_doc.add_table(rows=len(_sd)+1, cols=4); _t.style = "Table Grid"
for i, h in enumerate(["city", "n", "Moran's I", "permutation p"]):
    _t.rows[0].cells[i].text = h
for r, row in enumerate(_sd, 1):
    for i, k in enumerate(["city", "n", "moran_i", "perm_p_value"]):
        _t.rows[r].cells[i].text = str(row[k])
sup_doc.add_paragraph(
    "Supplementary Table S3. Zero-mass structure of VoF by zone: "
    "positive-VoF clusters per zone and conditional medians "
    "(CEUS_ZERO_MASS_SUMMARY.csv). Positive incidence differs across "
    "zones (3/360, 13/88, 8/106 of electrically feasible clusters; 2 further cases within numerical tolerance); "
    "the pooled ~96% zero share is driven mainly by Sano.")
_zm = list(_csv.DictReader(open(f"{ROOT}/analysis/CEUS_ZERO_MASS_SUMMARY.csv")))
_t2 = sup_doc.add_table(rows=len(_zm)+1, cols=5); _t2.style = "Table Grid"
for i, h in enumerate(["city", "feasible clusters", "positive VoF", "median VoF% | positive", "max VoF%"]):
    _t2.rows[0].cells[i].text = h
for r, row in enumerate(_zm, 1):
    for i, k in enumerate(["city", "eligible_feasible", "positive_vof", "median_vof_pct_given_positive", "max_vof_pct"]):
        _t2.rows[r].cells[i].text = str(row[k])
sup_doc.add_paragraph(
    "Morphology robustness note. Convex-hull demand density is "
    "degenerate for 427 of 554 feasible clusters (fewer than three terminals or "
    "hull area < 1e-4 km²); the hull-density association is therefore "
    "reported only on the 127 non-degenerate feasible clusters, and density is "
    "re-expressed on a fixed 100-m buffer support in the main text.")
sup_doc.add_paragraph(
    "Additional machine-readable tables (cluster-level results, "
    "morphology descriptors, engineering-assumption sensitivity, "
    "Monte-Carlo sweep) are provided as CSV files in the submission "
    "package.")
sup_doc.save(f"{ROOT}/manuscript/supplement_ceus.docx")

# ---------------- inline docx ----------------
doc = Document(f"{ROOT}/manuscript/manuscript_ceus_v1.docx")
pairs = [
    ("Figure 1 summarizes", ["ceus_fig_pipeline.png",
     "Figure 1. Data-to-decision pipeline: open geospatial inputs, zone construction, feasibility-gated reconstruction, and the incumbent-admissible renewal decision model."]),
    ("(Figure 2)", ["v4_fig1_study_zones.png",
     "Figure 2. Geographic context and study zones. (a) Locations of Sano (Tochigi), Kudamatsu (Yamaguchi) and Yasu (Shiga) within Japan (markers show municipal positions; analyzed areas are submunicipal; prefecture boundaries from geoBoundaries Japan ADM1 (William & Mary geoLab, https://www.geoboundaries.org/, CC-BY-4.0) and northern island coastlines from Natural Earth 10m land (public domain); markers distinguished in the legend). (b–d) Road graph (grey) and snapped building demand nodes (blue) in each selected study zone (Sano 772, Kudamatsu 341, Yasu 332 nodes). Municipal locations provide geographic context and are not representative of each city in its entirety."]),
    ("Figure 3 compares", ["v4_fig2_topology_example.png",
     "Figure 3. Optimized vs incumbent route length for all 554 feasible clusters (log-log; 19 clusters alter the corridor by >1 m)."]),
    ("(Figure 4)", ["v4_fig3_delta.png",
     "Figure 4. Primary estimand VoF by zone vs Monte-Carlo mean of the secondary naive ΔNPC; band = prespecified 1% practical-equivalence margin."]),
    ("(Figure 5, Table 3)", ["ceus_fig_morphology.png",
     "Figure 5. Feasible-cohort VoF mass structure: positive-VoF incidence by zone and conditional magnitudes for the 24 positive clusters (95.7% of feasible clusters have zero VoF)."]),
    ("(Figure 6)", ["ceus_fig_decomposition.png",
     "Figure 6. S1−S2 decomposition per zone in common monetary units (bn JPY, 40-yr discounted): capex, loss-cost and residual differences, alongside the non-negative VoF aggregate; panels use independent y-scales."]),
]
done = set()
for par in doc.paragraphs:
    for anchor_text, (fn, cap) in reversed(pairs):
        if anchor_text in par.text and anchor_text not in done:
            anchor = par._p
            pic = doc.add_paragraph()
            pic.add_run().add_picture(f"{ROOT}/figures/{fn}", width=Inches(5.8))
            capp = doc.add_paragraph(cap)
            anchor.addnext(capp._p); anchor.addnext(pic._p)
            done.add(anchor_text)
print("inserted:", done)
doc.save(f"{ROOT}/manuscript/manuscript_inline_ceus.docx")

# ---------------- readme ----------------
readme = """CEUS SUBMISSION README — portal upload map (Elsevier Editorial Manager)

1. Manuscript (main document): manuscript_ceus_final.docx
   (inline-figure convenience copy for review: manuscript_inline_ceus_final.docx)
2. Figures (separate files, both PNG and TIFF supplied):
   ceus_fig_pipeline, v4_fig1_study_zones, v4_fig2_topology_example,
   v4_fig3_delta, ceus_fig_morphology, ceus_fig_decomposition
   → order: pipeline=Fig1, study_zones=Fig2, topology_example=Fig3,
     delta=Fig4, morphology=Fig5, decomposition=Fig6
3. Supplementary material: supplement_ceus.docx (+ v4_supp_figS1_dg)
4. Highlights: highlights_ceus.txt (required by Elsevier)
5. Cover letter: cover_letter_ceus.txt
6. Declarations: declarations_ceus.txt
7. Data/code availability: data_code_availability_ceus.txt
8. Tables: embedded in main docx (Tables 1–4)
9. Fill author name, affiliation, and corresponding-author fields in
   the portal; declare AI use per Elsevier policy (text prepared in
   declarations).
10. Internal QA (not for journal upload): audit/, FINAL_HANDOFF.md,
    FINAL_SOURCE_MANIFEST.md.
"""
open(f"{ROOT}/manuscript/CEUS_SUBMISSION_README.txt", "w").write(readme)

# ---------------- package: complete manuscript/ subtree ----------------
pkg = f"{ROOT}/CEUS_manuscript_folder_FINAL"
if os.path.exists(pkg):
    shutil.rmtree(pkg)
os.makedirs(pkg)
sub_files = [
    "manuscript_ceus_final.docx", "manuscript_inline_ceus_final.docx",
    "supplement_ceus.docx", "cover_letter_ceus.txt", "highlights_ceus.txt",
    "declarations_ceus.txt", "data_code_availability_ceus.txt",
    "CEUS_SUBMISSION_README.txt", "FINAL_HANDOFF.md", "FINAL_SOURCE_MANIFEST.md",
]
os.makedirs(f"{pkg}/manuscript/figures", exist_ok=True)
os.makedirs(f"{pkg}/manuscript/tables", exist_ok=True)
for f in sub_files:
    shutil.copy(f"{ROOT}/manuscript/{f}", f"{pkg}/manuscript/{f}")
figset = ["ceus_fig_pipeline", "v4_fig1_study_zones", "v4_fig2_topology_example",
          "v4_fig3_delta", "ceus_fig_morphology", "ceus_fig_decomposition",
          "v4_supp_figS1_dg"]
for f in figset:
    for ext in ["png", "tiff"]:
        src = f"{ROOT}/figures/{f}.{ext}"
        if os.path.exists(src):
            shutil.copy(src, f"{pkg}/manuscript/figures/{f}.{ext}")
tables = ["analysis/CEUS_MORPHOLOGY_CLUSTER.csv",
          "analysis/CEUS_MORPHOLOGY_CORRELATIONS.csv",
          "analysis/CEUS_BOUNDARY_SENSITIVITY.csv",
          "analysis/CEUS_ZERO_MASS_SUMMARY.csv",
          "analysis/CEUS_SPATIAL_DIAGNOSTICS.csv",
          "analysis/scenario_results_t130.csv",
          "analysis/manuscript_values_v3.csv", "analysis/scenario_results_v2.csv"]
for f in tables:
    shutil.copy(f"{ROOT}/{f}", f"{pkg}/manuscript/tables/{os.path.basename(f)}")
for a in glob.glob(f"{ROOT}/manuscript/audit/*"):
    os.makedirs(f"{pkg}/manuscript/audit", exist_ok=True)
    shutil.copy(a, f"{pkg}/manuscript/audit/{os.path.basename(a)}")
zpath = f"{ROOT}/CEUS_manuscript_folder_SUBMISSION_LOCKED.zip"
if os.path.exists(zpath):
    os.remove(zpath)
with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
    for dp, _, fs in os.walk(pkg):
        for f in fs:
            fp = os.path.join(dp, f)
            z.write(fp, os.path.relpath(fp, pkg))
print("zip entries:", len(zipfile.ZipFile(zpath).namelist()))
print("done")
