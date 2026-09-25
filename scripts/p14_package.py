"""Build an allowlisted, contamination-scanned SEGAN submission package."""
import csv
import datetime
import hashlib
import json
import os
import platform
import re
import subprocess
import zipfile

import pandas as pd
from docx import Document

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
NOW = datetime.datetime.now(datetime.timezone.utc).isoformat()
PYTHON = f"{ROOT}/.venv/bin/python"
values = {
    row["key"]: row["value"]
    for row in csv.DictReader(open(
        f"{ROOT}/FINAL_MANUSCRIPT_VALUES.csv", encoding="utf-8"))
}
scenario = pd.read_csv(
    f"{ROOT}/analysis/scenario_results.csv").set_index("city")
sensitivity = pd.read_csv(
    f"{ROOT}/analysis/sensitivity_mc.csv").set_index("city")


def v(key):
    return float(values[key])


def f(key, digits=1):
    return f"{v(key):.{digits}f}"


os.makedirs(f"{ROOT}/manuscript", exist_ok=True)
cover = f"""Dear Editor,

Please consider “Engineering-feasible screening of underground distribution
renewal: separating routing-policy and conductor-sizing effects” as an
Original Research Article in Sustainable Energy, Grids and Networks.

The manuscript presents a reproducible forensic reanalysis with an explicit
MV/transformer/LV hierarchy, active/reactive loss propagation, quantitative AC
benchmark comparison and binding feasibility gates. Fixed-conductor routing-
policy savings are {f('sano_routing_policy_savings_pct')}%,
{f('kudamatsu_routing_policy_savings_pct')}% and
{f('yasu_routing_policy_savings_pct')}%; combined best-found routing and
conductor-screening savings are {f('sano_combined_screening_savings_pct')}%,
{f('kudamatsu_combined_screening_savings_pct')}% and
{f('yasu_combined_screening_savings_pct')}%. The cases are explicitly framed
as synthetic geospatial screens rather than observed municipal utility
networks or construction-ready designs.

The work is original, is not under consideration elsewhere, and has been
approved by all authors. Author, affiliation and corresponding-author fields
will be completed locally before submission.

Sincerely,
[Corresponding author]
"""
with open(
        f"{ROOT}/manuscript/cover_letter_corrected.txt", "w",
        encoding="utf-8") as handle:
    handle.write(cover)

declarations = """Funding: None.
Declaration of competing interests: None declared.
Ethics approval: Not applicable; no human participants or personal data.
Data and code availability: Raw public-data snapshots, retrieval metadata,
checksums, scripts, machine-readable outputs and reproduction instructions
accompany the reproducibility package.
Generative AI: During preparation, the authors used an AI coding assistant for
code review, reproducibility checking and language drafting. The authors
reviewed and edited all outputs and take responsibility for the content.
"""
with open(
        f"{ROOT}/manuscript/declarations_corrected.txt", "w",
        encoding="utf-8") as handle:
    handle.write(declarations)

os.makedirs(f"{ROOT}/municipality", exist_ok=True)
for city, japanese_name in [
        ("sano", "佐野市"), ("kudamatsu", "下松市"), ("yasu", "野洲市")]:
    summary = f"""# {japanese_name} planning-screening summary

This is a synthetic geospatial research scenario, not an engineering design or
a reconstruction of the municipality's observed utility network.

- Modeled demand: {v(city + '_total_demand_kw')/1000:.2f} MW
- S1 static-horizon NPC: {v(city + '_s1_npc')/1e9:.2f} billion JPY
- S1.5 static-horizon NPC: {v(city + '_s15_npc')/1e9:.2f} billion JPY
- S2 best-found NPC: {v(city + '_s2_npc')/1e9:.2f} billion JPY
- Fixed-conductor routing-policy savings:
  {v(city + '_routing_policy_savings_pct'):.2f}%
- Combined routing-and-conductor screening savings:
  {v(city + '_combined_screening_savings_pct'):.2f}%
- Maximum modeled MV drop: {v(city + '_max_mv_vdrop_pct'):.3f}%
- Minimum equivalent customer voltage:
  {v(city + '_min_customer_voltage_pu'):.4f} pu
- Conditional positive-draw frequency:
  {v(city + '_sampling_positive_draw_frequency'):.3f}

Topology and equipment are frozen within each scenario-sampling draw. Project
decisions require utility topology, measured load, shared-substation modeling,
field survey, easement, reliability, protection and local bid data.
"""
    with open(
            f"{ROOT}/municipality/{city}_planning_support_corrected.md", "w",
            encoding="utf-8") as handle:
        handle.write(summary)

claims = [
    (
        "All reported clusters pass the frozen screening gates",
        "analysis/scenario_results.csv;analysis/feasibility_diagnostics.csv",
        "modeled"),
    (
        "Fixed-conductor routing-policy and combined screening savings are "
        "distinct estimands",
        "analysis/scenario_results.csv", "modeled"),
    (
        "Paired route/conductor attribution is policy-specific and non-causal",
        "analysis/ablation_decomposition.csv", "modeled"),
    (
        "Three CIGRE radial feeders pass predeclared error tolerances",
        "analysis/benchmark_radial_comparison.csv", "benchmark"),
    (
        "case11_iwamoto is excluded after three solver failures",
        "analysis/benchmark_ac_solvers.csv;"
        "analysis/benchmark_iwamoto_diagnosis.csv", "benchmark"),
    (
        "Municipal datasets do not observe feeder connectivity",
        "MUNICIPAL_DATA_PROVENANCE_AUDIT.csv", "observed/provenance"),
    (
        "Cost inputs are assumptions varied in conditional scenario sampling",
        "COST_PARAMETER_PROVENANCE.csv;analysis/sensitivity_mc.csv",
        "assumed/sensitivity"),
    (
        "Topology and equipment are frozen within each sampling draw",
        "analysis/sensitivity_mc.csv", "sensitivity"),
    (
        "Prefix stability is not independent-seed algorithmic convergence",
        "analysis/reconstruction_prefix_stability_summary.csv", "diagnostic"),
]
pd.DataFrame(
    claims, columns=["claim", "evidence_source", "epistemic_label"]).to_csv(
        f"{ROOT}/CLAIM_EVIDENCE_MATRIX.csv", index=False)

requirements = subprocess.run(
    [PYTHON, "-m", "pip", "freeze"], text=True, capture_output=True,
    check=True).stdout
environment = {
    "generated_utc": NOW,
    "python": platform.python_version(),
    "platform": platform.platform(),
    "packages": requirements.splitlines(),
}
with open(f"{ROOT}/ENVIRONMENT_SNAPSHOT.json", "w", encoding="utf-8") as handle:
    json.dump(environment, handle, indent=2)

reproduction = f"""# Corrected SEGAN reproduction

The reference build used Python {platform.python_version()}. Create an isolated
environment with that Python version and install `requirements.lock.txt`.
Archived public inputs are under `data/raw/`; canonical checksums are recorded
under `audit/`.

Run:

```bash
make final
```

The build regenerates provenance, catalogues, audits, the benchmark,
geospatial demand, reconstruction, feasibility-gated optimization,
conditional scenario sampling, the frozen value ledger, figures, manuscripts,
tests, inventories and the allowlisted ZIP. Runtime is dominated by
reconstruction candidates and bounded conductor search.
"""
with open(
        f"{ROOT}/REPRODUCTION_CORRECTED.md", "w", encoding="utf-8") as handle:
    handle.write(reproduction)

manifest = pd.read_csv(f"{ROOT}/audit/ORIGINALS_MANIFEST.csv")
original_checks = []
for _, row in manifest.iterrows():
    relative_path = f"audit/originals/{row['filename']}"
    path = f"{ROOT}/{relative_path}"
    digest = hashlib.sha256(open(path, "rb").read()).hexdigest()
    original_checks.append({
        "path": relative_path,
        "expected_sha256": row["expected_sha256"],
        "actual_sha256": digest,
        "pass": digest == row["expected_sha256"],
    })

fixed_files = [
    "ENGINEERING_FEASIBILITY_SPEC.yaml",
    "VOLTAGE_ROOT_CAUSE_AUDIT.md",
    "BENCHMARK_VALIDATION_REPORT.md",
    "ORIGINAL_VS_CORRECTED_RESULTS.csv",
    "MUNICIPAL_DATA_PROVENANCE_AUDIT.csv",
    "COST_PARAMETER_PROVENANCE.csv",
    "RAW_DATA_ACQUISITION_REGISTRY.csv",
    "FINAL_MANUSCRIPT_VALUES.csv",
    "CLAIM_EVIDENCE_MATRIX.csv",
    "ENVIRONMENT_SNAPSHOT.json",
    "REPRODUCTION_CORRECTED.md",
    "requirements.lock.txt",
    "Makefile",
    "config/study.yaml",
    "AUDIT_FILE_INVENTORY.csv",
    "AUDIT_DEPENDENCY_DAG.csv",
    "AUDIT_DEPENDENCY_DAG.md",
    "ENVIRONMENT_SOLVER_VERSIONS.csv",
    "DECISION_LOG.md",
    "PHASE_0_QC.json",
    "PHASE_0_HANDOFF.txt",
    "PHASE_1_QC.json",
    "PHASE_1_HANDOFF.txt",
    "audit/ORIGINALS_MANIFEST.csv",
    "audit/ORIGINAL_NUMERIC_TEXT_CONTEXTS.csv",
    "audit/README.md",
    "audit/originals/segan_submission_files.zip",
    "audit/originals/manuscript_inline.docx",
    "audit/original_results/scenario_results.csv",
    "audit/CLEAN_ROOM_REPRODUCTION.md",
    "provenance/decision_log.md",
    "provenance/municipality_candidates.csv",
    "provenance/source_registry.csv",
    "data/processed/engineering/cable_catalog.csv",
    "data/processed/engineering/transformer_catalog.csv",
    "analysis/ablation_decomposition.csv",
    "analysis/ablation_results.csv",
    "analysis/benchmark_ac_solvers.csv",
    "analysis/benchmark_iwamoto_diagnosis.csv",
    "analysis/benchmark_radial_comparison.csv",
    "analysis/benchmark_validation.json",
    "analysis/candidate_search_summary.csv",
    "analysis/component_decomposition.csv",
    "analysis/data_audit.csv",
    "analysis/feasibility_diagnostics.csv",
    "analysis/manuscript_values.csv",
    "analysis/reconstruction_prefix_stability.csv",
    "analysis/reconstruction_prefix_stability_summary.csv",
    "analysis/reconstruction_summary.csv",
    "analysis/references_verified.csv",
    "analysis/scenario_results.csv",
    "analysis/sensitivity_draws.csv",
    "analysis/sensitivity_mc.csv",
    "analysis/sensitivity_scenario_sampling.csv",
    "analysis/sensitivity_tornado.csv",
    "analysis/transformer_service_units.csv",
    "figures/fig1_study_zones.png",
    "figures/fig1_study_zones.tiff",
    "figures/fig2_topology_example.png",
    "figures/fig2_topology_example.tiff",
    "figures/fig3_savings_uncertainty.png",
    "figures/fig3_savings_uncertainty.tiff",
    "figures/fig4_ablation_decomposition.png",
    "figures/fig4_ablation_decomposition.tiff",
    "tables/table2_corrected_results.csv",
    "tables/table3_uncertainty.csv",
    "tables/table4_ablation.csv",
    "manuscript/manuscript_corrected.docx",
    "manuscript/manuscript_inline_corrected.docx",
    "manuscript/supplementary_material_corrected.docx",
    "manuscript/cover_letter_corrected.txt",
    "manuscript/declarations_corrected.txt",
    "manuscript/highlights_corrected.txt",
    "municipality/sano_planning_support_corrected.md",
    "municipality/kudamatsu_planning_support_corrected.md",
    "municipality/yasu_planning_support_corrected.md",
]
script_files = [
    "scripts/electrical_model.py",
    "scripts/p00_forensic_inventory.py",
    "scripts/p01_traceability.py",
    "scripts/p03_data_audit.py",
    "scripts/p04_benchmark.py",
    "scripts/p05_gis_demand.py",
    "scripts/p06_reconstruct.py",
    "scripts/p07_optimize.py",
    "scripts/p10_sensitivity.py",
    "scripts/p11_freeze.py",
    "scripts/p12_figures.py",
    "scripts/p13_manuscript.py",
    "scripts/p14_package.py",
    "scripts/p15_inline_docx.py",
    "scripts/p16_catalogs.py",
    "scripts/p17_results_audit.py",
    "scripts/p18_provenance_convergence.py",
    "scripts/p19_raw_provenance.py",
    "tests/test_electrical_model.py",
]
raw_registry = pd.read_csv(f"{ROOT}/RAW_DATA_ACQUISITION_REGISTRY.csv")
raw_files = sorted(set(raw_registry.local_path.astype(str)))
allowlist = sorted(set(fixed_files + script_files + raw_files))

missing = [path for path in allowlist if not os.path.isfile(f"{ROOT}/{path}")]
if missing:
    raise RuntimeError(f"allowlisted files missing: {missing}")

forbidden_parts = [
    ("topology " + "likelihood distance"),
    ("roughly " + "halved"),
    ("P(" + "LTP>0)"),
    ("life-cycle " + "topology potential"),
]
legacy_numbers = [
    "12.1956", "2.5441", "11.7190", "29.1596", "31.1232", "37.7830",
    "61.7%", "60.1%", "65.0%",
]
forensic_contexts = {
    "ORIGINAL_VS_CORRECTED_RESULTS.csv",
    "VOLTAGE_ROOT_CAUSE_AUDIT.md",
    "audit/ORIGINAL_NUMERIC_TEXT_CONTEXTS.csv",
    "audit/originals/manuscript_inline.docx",
    "DECISION_LOG.md",
    "scripts/p01_traceability.py",
    "scripts/p17_results_audit.py",
    "scripts/p14_package.py",
}


def docx_text(path):
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_text(relative_path):
    path = f"{ROOT}/{relative_path}"
    if relative_path.endswith(".docx"):
        return docx_text(path)
    if os.path.basename(relative_path) == "Makefile" or relative_path.endswith(
            (".txt", ".md", ".csv", ".json", ".yaml", ".yml", ".py")):
        return open(path, encoding="utf-8", errors="replace").read()
    return None


scan_rows = []
for relative_path in allowlist:
    text = extract_text(relative_path)
    if text is None:
        continue
    lowered = text.lower()
    for pattern in forbidden_parts:
        if pattern.lower() in lowered:
            scan_rows.append({
                "path": relative_path, "pattern": pattern,
                "classification": "forensic_context"
                if relative_path in forensic_contexts else "prohibited",
            })
    for pattern in legacy_numbers:
        if re.search(rf"(?<!\d){re.escape(pattern)}(?!\d)", text):
            scan_rows.append({
                "path": relative_path, "pattern": pattern,
                "classification": "forensic_context"
                if relative_path in forensic_contexts else "prohibited",
            })
scan = pd.DataFrame(
    scan_rows, columns=["path", "pattern", "classification"])
scan.to_csv(f"{ROOT}/PACKAGE_TEXT_SCAN.csv", index=False)
allowlist.extend(["PACKAGE_TEXT_SCAN.csv"])
allowlist = sorted(set(allowlist))

prohibited_hits = (
    scan[scan.classification == "prohibited"].to_dict("records")
    if not scan.empty else [])
qc = {
    "generated_utc": NOW,
    "canonical_originals": original_checks,
    "all_canonical_originals_unchanged": all(
        item["pass"] for item in original_checks),
    "all_reported_networks_feasible": bool(
        scenario.all_primary_networks_feasible.all()),
    "scenario_sampling_draws_per_city": {
        city: int(sensitivity.loc[city, "n_draws"])
        for city in scenario.index
    },
    "coherent_city_reconstruction_draw": bool(
        sensitivity.coherent_city_reconstruction_draw.all()),
    "topology_reoptimized_each_draw": bool(
        sensitivity.topology_reoptimized_each_draw.any()),
    "equipment_resized_each_draw": bool(
        sensitivity.equipment_resized_each_draw.any()),
    "package_text_files_scanned": int(
        sum(extract_text(path) is not None for path in allowlist
            if os.path.exists(f"{ROOT}/{path}"))),
    "prohibited_active_package_hits": prohibited_hits,
    "forensic_context_hits": (
        scan[scan.classification == "forensic_context"].to_dict("records")
        if not scan.empty else []),
}
qc["status"] = "pass" if (
    qc["all_canonical_originals_unchanged"]
    and qc["all_reported_networks_feasible"]
    and qc["coherent_city_reconstruction_draw"]
    and not qc["topology_reoptimized_each_draw"]
    and not qc["equipment_resized_each_draw"]
    and not prohibited_hits
) else "fail"
with open(f"{ROOT}/FINAL_QC_MACHINE.json", "w", encoding="utf-8") as handle:
    json.dump(qc, handle, indent=2)

report = f"""# Final machine QC

- Canonical originals unchanged: {qc['all_canonical_originals_unchanged']}
- All reported networks feasible: {qc['all_reported_networks_feasible']}
- Coherent city-level reconstruction sampling:
  {qc['coherent_city_reconstruction_draw']}
- Topology reoptimized within draws: {qc['topology_reoptimized_each_draw']}
- Equipment resized within draws: {qc['equipment_resized_each_draw']}
- Text-extractable package files scanned:
  {qc['package_text_files_scanned']}
- Prohibited active-package hits: {len(prohibited_hits)}
- Machine verdict: {qc['status']}

Forensic files may contain quarantined legacy values when explicitly labeled.
Clean-room and hostile-review records are separate mandatory release gates.
"""
with open(f"{ROOT}/FINAL_QC_REPORT.md", "w", encoding="utf-8") as handle:
    handle.write(report)

handoff = f"""SEGAN corrected-package handoff
Generated: {NOW}
Routing-policy savings: {f('sano_routing_policy_savings_pct')}% /
{f('kudamatsu_routing_policy_savings_pct')}% /
{f('yasu_routing_policy_savings_pct')}%.
Combined screening savings: {f('sano_combined_screening_savings_pct')}% /
{f('kudamatsu_combined_screening_savings_pct')}% /
{f('yasu_combined_screening_savings_pct')}%.
These are synthetic, assumption-conditional screening results.
"""
with open(f"{ROOT}/FINAL_HANDOFF.txt", "w", encoding="utf-8") as handle:
    handle.write(handoff)

upload_set = """SEGAN upload set
1. manuscript/manuscript_corrected.docx
2. manuscript/cover_letter_corrected.txt
3. manuscript/highlights_corrected.txt
4. manuscript/declarations_corrected.txt
5. manuscript/supplementary_material_corrected.docx
6. figures/fig1_study_zones.tiff
7. figures/fig2_topology_example.tiff
8. figures/fig3_savings_uncertainty.tiff
9. figures/fig4_ablation_decomposition.tiff

The inline manuscript is a review copy, not a second submission manuscript.
"""
with open(
        f"{ROOT}/SUBMISSION_UPLOAD_SET.txt", "w", encoding="utf-8") as handle:
    handle.write(upload_set)

for path in [
        "FINAL_QC_MACHINE.json", "FINAL_QC_REPORT.md", "FINAL_HANDOFF.txt",
        "SUBMISSION_UPLOAD_SET.txt", "PACKAGE_ALLOWLIST.json",
        "PACKAGE_FILE_INVENTORY.csv",
]:
    if path not in allowlist:
        allowlist.append(path)
allowlist = sorted(set(allowlist))

with open(f"{ROOT}/PACKAGE_ALLOWLIST.json", "w", encoding="utf-8") as handle:
    json.dump({"generated_utc": NOW, "files": allowlist}, handle, indent=2)

inventory = []
for relative_path in allowlist:
    if relative_path == "PACKAGE_FILE_INVENTORY.csv":
        continue
    path = f"{ROOT}/{relative_path}"
    data = open(path, "rb").read()
    inventory.append({
        "path": relative_path,
        "size_bytes": len(data),
        "sha256": hashlib.sha256(data).hexdigest(),
    })
pd.DataFrame(inventory).to_csv(f"{ROOT}/PACKAGE_FILE_INVENTORY.csv", index=False)

zip_path = f"{ROOT}/SEGAN_submission_package_FINAL.zip"
with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
    for relative_path in allowlist:
        archive.write(f"{ROOT}/{relative_path}", relative_path)
with zipfile.ZipFile(zip_path) as archive:
    corrupt = archive.testzip()
    archived = sorted(archive.namelist())
if corrupt is not None:
    raise RuntimeError(f"ZIP CRC failure: {corrupt}")
if archived != sorted(allowlist):
    raise RuntimeError("ZIP contents differ from package allowlist")
if qc["status"] != "pass":
    raise RuntimeError("final machine QC failed")
print(
    "QC", qc["status"], "ZIP", zip_path,
    f"{os.path.getsize(zip_path)/1e6:.1f} MB",
    "files", len(allowlist))
