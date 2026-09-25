"""Insert corrected figures into a review-copy DOCX."""
import os

from docx import Document
from docx.shared import Inches

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
doc = Document(f"{ROOT}/manuscript/manuscript_corrected.docx")
insertions = [
    ("(Figure 1; Table 1)", ["fig1_study_zones"]),
    ("Figure 2 shows", ["fig2_topology_example"]),
    ("(Table 2; Figure 3)", ["fig3_savings_uncertainty"]),
    ("(Figure 4; Table 3)", ["fig4_ablation_decomposition"]),
]
captions = {
    "fig1_study_zones": "Figure 1. Study zones and modeled demand nodes.",
    "fig2_topology_example": (
        "Figure 2. Example reconstructed and best-found candidate routes."),
    "fig3_savings_uncertainty": (
        "Figure 3. Base-case and conditional scenario-sampling savings."),
    "fig4_ablation_decomposition": (
        "Figure 4. Paired Shapley attribution to routing and conductors."),
}


def insert_after(paragraph, figure_names):
    anchor = paragraph._p
    for name in figure_names:
        picture = doc.add_paragraph()
        picture.add_run().add_picture(
            f"{ROOT}/figures/{name}.png", width=Inches(5.8))
        caption = doc.add_paragraph(captions[name])
        anchor.addnext(caption._p)
        anchor.addnext(picture._p)
        anchor = caption._p


done = set()
for paragraph in list(doc.paragraphs):
    for anchor_text, figures in insertions:
        if anchor_text in paragraph.text and anchor_text not in done:
            insert_after(paragraph, figures)
            done.add(anchor_text)
if len(done) != len(insertions):
    raise RuntimeError(f"missing insertion anchors: {set(x[0] for x in insertions)-done}")
path = f"{ROOT}/manuscript/manuscript_inline_corrected.docx"
doc.save(path)
print("saved", path)
