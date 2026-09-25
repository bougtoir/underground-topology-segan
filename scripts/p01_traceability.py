"""Audit every quantitative manuscript claim against code and raw dependencies."""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path

import pandas as pd
import yaml
from docx import Document


ROOT = Path(__file__).resolve().parents[1]


def claim(
    claim_id: str,
    location: str,
    text: str,
    value: str,
    unit: str,
    source_file: str,
    source_variable: str,
    code: str,
    raw_dependency: str,
    status: str,
    severity: str,
    note: str,
) -> dict[str, str]:
    return locals()


def append_decision(text: str) -> None:
    path = ROOT / "DECISION_LOG.md"
    existing = path.read_text(encoding="utf-8")
    if text not in existing:
        with path.open("a", encoding="utf-8") as handle:
            handle.write(text + "\n")


def manuscript_text() -> list[tuple[str, str]]:
    doc = Document(ROOT / "audit/originals/manuscript_inline.docx")
    rows = []
    section = "front matter"
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style.name.startswith("Heading") or text in {"Abstract", "References", "Declarations"}:
            section = text
        rows.append((f"{section}:paragraph-{index + 1}", text))
    for table_index, table in enumerate(doc.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            rows.append(
                (
                    f"table-{table_index}:row-{row_index}",
                    " | ".join(cell.text for cell in row.cells),
                )
            )
    return rows


def main() -> None:
    scen = pd.read_csv(ROOT / "analysis/scenario_results.csv").set_index("city")
    mc = pd.read_csv(ROOT / "analysis/sensitivity_mc.csv").set_index("city")
    dg = pd.read_csv(ROOT / "analysis/dg_results.csv").set_index("city")
    data = pd.read_csv(ROOT / "analysis/data_audit.csv").set_index("municipality")
    cfg = yaml.safe_load((ROOT / "config/study.yaml").read_text(encoding="utf-8"))

    rows = []
    total_clusters = int(scen["clusters"].sum())
    total_demand = float(scen["total_demand_kw"].sum() / 1000)
    rows.extend(
        [
            claim("Q001", "Abstract; Results", "Total modeled feeder clusters", str(total_clusters), "count",
                  "analysis/scenario_results.csv", "sum(clusters)", "scripts/p07_optimize.py",
                  "processed demand nodes and road graph", "reproducible-derived", "MINOR",
                  "Derived sum should be frozen explicitly in the corrected value registry."),
            claim("Q002", "Abstract; Results", "Total modeled peak demand", f"{total_demand:.1f}", "MW",
                  "analysis/scenario_results.csv", "sum(total_demand_kw)/1000", "scripts/p07_optimize.py",
                  "OSM building footprints; modeled demand densities", "reproducible-derived", "MINOR",
                  "The quantity is modeled, not observed."),
        ]
    )
    next_id = 3
    for city in ("sano", "kudamatsu", "yasu"):
        display = city.capitalize()
        values = [
            ("LTP percentage", scen.loc[city, "ltp_pct"], "%", "ltp_pct", "MAJOR"),
            ("NPC difference", scen.loc[city, "delta_npc"] / 1e9, "bn JPY", "delta_npc", "MAJOR"),
            ("Conductor upsizing distance", scen.loc[city, "upsize_km"], "km", "upsize_km", "MAJOR"),
            ("Peak-loss reduction", 100 * (1 - scen.loc[city, "s2_plkw"] / scen.loc[city, "s1_plkw"]), "%", "derived loss reduction", "MAJOR"),
            ("Maximum voltage drop", scen.loc[city, "max_vdrop"], "%", "max_vdrop", "CRITICAL"),
            ("Mean edge-set distance", scen.loc[city, "mean_tld"], "fraction", "mean_tld", "MAJOR"),
            ("Cluster-level positive value fraction", scen.loc[city, "p_ltp_pos"], "fraction", "p_ltp_pos", "MAJOR"),
            ("Monte Carlo probability of positive value", mc.loc[city, "p_ltp_pos"], "probability", "p_ltp_pos", "MAJOR"),
            ("Modeled rooftop PV capacity", dg.loc[city, "pv_kwp_total"] / 1000, "MWp", "pv_kwp_total", "MAJOR"),
        ]
        for label, value, unit, variable, severity in values:
            status = "reproducible-but-invalid-primary-model" if severity == "CRITICAL" else "reproducible-original-model"
            note = (
                "Voltage infeasibility invalidates primary scenario inclusion and the same-model argument."
                if severity == "CRITICAL"
                else "Must be regenerated after the feasibility repair."
            )
            rows.append(
                claim(
                    f"Q{next_id:03d}",
                    "Abstract; Results; Discussion; tables",
                    f"{display}: {label}",
                    f"{float(value):.12g}",
                    unit,
                    "analysis/scenario_results.csv" if label not in {"Monte Carlo probability of positive value", "Modeled rooftop PV capacity"} else (
                        "analysis/sensitivity_mc.csv" if "Monte Carlo" in label else "analysis/dg_results.csv"
                    ),
                    variable,
                    "scripts/p07_optimize.py; scripts/p10_sensitivity.py; scripts/p08_dg.py",
                    "OSM inputs; modeled demand; reconstruction ensemble; config/study.yaml",
                    status,
                    severity,
                    note,
                )
            )
            next_id += 1

    rows.extend(
        [
            claim(f"Q{next_id:03d}", "Methods", "Sano municipal record count", str(int(data.loc["Sano", "rows"])), "rows",
                  "analysis/data_audit.csv", "rows", "scripts/p03_data_audit.py",
                  "data/raw/sano/sano_security_lights.csv", "reproducible", "MINOR",
                  "Pole-ID semantics and license require source-level verification."),
            claim(f"Q{next_id + 1:03d}", "Methods", "Sano duplicated pole-field values", str(int(data.loc["Sano", "pole_dupes"])), "count",
                  "analysis/data_audit.csv", "pole_dupes", "scripts/p03_data_audit.py",
                  "data/raw/sano/sano_security_lights.csv", "reproducible", "MAJOR",
                  "duplicated() counts repeated rows, not distinct multi-light poles."),
            claim(f"Q{next_id + 2:03d}", "Methods", "Kudamatsu municipal record count", str(int(data.loc["Kudamatsu", "rows"])), "rows",
                  "analysis/data_audit.csv", "rows", "scripts/p03_data_audit.py",
                  "three Kudamatsu municipal CSVs", "reproducible", "MINOR",
                  "License and dataset completeness require source verification."),
            claim(f"Q{next_id + 3:03d}", "Methods", "Nominal cluster target", "350", "kW",
                  "scripts/p06_reconstruct.py", "CLUSTER_KW", "scripts/p06_reconstruct.py",
                  "assumption only", "hard-coded-unfrozen", "CRITICAL",
                  "This structural assumption produces infeasible LV feeders and is not in the frozen configuration."),
            claim(f"Q{next_id + 4:03d}", "Methods", "Reconstruction replicate count", str(cfg["reconstruction"]["n_replicates"]), "count",
                  "config/study.yaml", "reconstruction.n_replicates", "scripts/p06_reconstruct.py",
                  "assumption only", "config-backed", "MAJOR",
                  "Twelve replicates have no convergence evidence."),
            claim(f"Q{next_id + 5:03d}", "Methods", "LV conductor options", "150;250;400;600", "mm2",
                  "scripts/p07_optimize.py", "CONDUCTOR_SIZES", "scripts/p07_optimize.py",
                  "assumption only", "hard-coded-unfrozen", "MAJOR",
                  "Sizes, ampacities, reactances, and cost provenance are absent."),
            claim(f"Q{next_id + 6:03d}", "Methods", "Usable rooftop fraction", str(cfg["dg"]["rooftop_fraction_usable"]), "fraction",
                  "config/study.yaml", "dg.rooftop_fraction_usable", "scripts/p08_dg.py",
                  "OSM building footprints plus assumption", "config-backed-weak-evidence", "MAJOR",
                  "Requires provenance and sensitivity; S3 must be secondary."),
            claim(f"Q{next_id + 7:03d}", "Methods", "PV density", "0.15", "kWp/m2",
                  "scripts/p08_dg.py", "PV_KW_M2", "scripts/p08_dg.py",
                  "assumption only", "hard-coded-unfrozen", "MAJOR",
                  "No source or uncertainty range is recorded."),
            claim(f"Q{next_id + 8:03d}", "Methods", "PV loss-reduction cap", "50", "%",
                  "scripts/p08_dg.py", "loss_red_frac cap", "scripts/p08_dg.py",
                  "assumption only", "hard-coded-unfrozen", "MAJOR",
                  "The cap is not a power-flow or reverse-flow constraint."),
            claim(f"Q{next_id + 9:03d}", "Sensitivity", "Monte Carlo draws", str(int(cfg["sensitivity"]["n_mc"])), "draws",
                  "config/study.yaml", "sensitivity.n_mc", "scripts/p10_sensitivity.py",
                  "scenario/cost assumptions", "config-backed-insufficient-resolution", "MAJOR",
                  "No convergence study, material correlation model, or topology re-optimization."),
            claim(f"Q{next_id + 10:03d}", "Sensitivity", "Discount-rate range", "2-4.5", "%",
                  "scripts/p10_sensitivity.py", "draws.R", "scripts/p10_sensitivity.py",
                  "assumption only", "hard-coded-unfrozen", "MAJOR",
                  "Range differs from the config comment and lacks cost-model provenance."),
            claim(f"Q{next_id + 11:03d}", "Sensitivity", "Electricity-value range", "15-35", "JPY/kWh",
                  "scripts/p10_sensitivity.py", "draws.EKWH", "scripts/p10_sensitivity.py",
                  "assumption only", "hard-coded-unfrozen", "MAJOR",
                  "Range differs from the config comment and lacks a price-year source."),
            claim(f"Q{next_id + 12:03d}", "Discussion", "CIGRE LV model validation claim", "validated", "claim",
                  "analysis/benchmark_validation.csv", "converged", "scripts/p04_benchmark.py",
                  "pandapower bundled network", "unsupported", "CRITICAL",
                  "Convergence alone is not quantitative validation; the stored 0.912 pu minimum is outside the stated 0.95 pu gate."),
            claim(f"Q{next_id + 13:03d}", "Discussion", "Iwamoto 11-bus convergence", "false", "boolean",
                  "analysis/benchmark_validation.csv", "converged", "scripts/p04_benchmark.py",
                  "pandapower bundled network", "reproducible-but-undiagnosed", "MAJOR",
                  "Failure cause and suitability were not investigated."),
        ]
    )

    doc_rows = manuscript_text()
    numeric_contexts = []
    in_references = False
    for location, text in doc_rows:
        if text == "References":
            in_references = True
        if in_references:
            continue
        if re.search(r"\d", text):
            numeric_contexts.append({"location": location, "text": text})
    with (ROOT / "audit/ORIGINAL_NUMERIC_TEXT_CONTEXTS.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=["location", "text"], lineterminator="\n")
        writer.writeheader()
        writer.writerows(numeric_contexts)

    out = ROOT / "MANUSCRIPT_VALUE_AUDIT.csv"
    with out.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=rows[0].keys(), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)

    severity_counts = pd.Series([row["severity"] for row in rows]).value_counts().to_dict()
    critical = [row for row in rows if row["severity"] == "CRITICAL"]
    hard_coded = [row for row in rows if "hard-coded" in row["status"]]
    qc = {
        "phase": 1,
        "status": "reanalysis_required",
        "audited_claims": len(rows),
        "numeric_text_contexts": len(numeric_contexts),
        "severity_counts": severity_counts,
        "hard_coded_quantities": len(hard_coded),
        "critical_findings": [row["note"] for row in critical],
    }
    (ROOT / "PHASE_1_QC.json").write_text(json.dumps(qc, indent=2) + "\n", encoding="utf-8")
    (ROOT / "PHASE_1_HANDOFF.txt").write_text(
        "Phase 1 traced the original manuscript's quantitative claims to code and raw "
        f"dependencies. {len(critical)} CRITICAL findings require reanalysis: electrically "
        "infeasible primary feeders, unsupported benchmark validation, and the unfrozen "
        "350-kW LV clustering assumption. Original headline values are quarantined from "
        "the corrected manuscript until regenerated.\n",
        encoding="utf-8",
    )
    append_decision(
        "2026-09-24 | Phase 1 | Quarantine all original headline effect sizes from the "
        "revised manuscript until feasibility-constrained regeneration | The 29-38% voltage "
        "drops are a fatal primary-model failure; using the same invalid model for both "
        "scenarios does not establish feasibility or comparability. | frozen"
    )
    append_decision(
        "2026-09-24 | Phase 1 | Rename topology likelihood distance in the corrected work | "
        "The implemented quantity is a Jaccard edge-set distance and contains no likelihood "
        "model. | frozen"
    )
    print(json.dumps(qc, indent=2))


if __name__ == "__main__":
    main()
